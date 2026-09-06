import os
import re

import openpyxl
from PySide6.QtCore import QThread, Signal
# 重评逻辑

class ReReviewWorker(QThread):
    """Background worker for a re-review task."""

    progress = Signal(int)
    finished = Signal(tuple)
    error = Signal(str)

    def __init__(self, excel1_path, excel2_path, output_dir, project_name=None):
        super().__init__()
        self.excel1_path = excel1_path
        self.excel2_path = excel2_path
        self.output_dir = output_dir
        self.project_name = project_name or "ReReview"

    def run(self):
        from utils.runtime_logger import RuntimeLogger

        RuntimeLogger.set_project(self.project_name)
        RuntimeLogger.log("开始重评处理...")
        try:
            result = ReReviewProcessor.process_re_review(
                self.excel1_path, self.excel2_path, self.output_dir, self.progress.emit
            )
            self.finished.emit(result)
        except Exception as exc:
            import traceback

            traceback.print_exc()
            self.error.emit(str(exc))


class ReReviewProcessor:
    """Create re-review annotations between the old report and new receipt."""

    START_ROW = 5
    CURRENT_REUSE_COL = 20  # T
    NEW_ID_COL = 21  # U
    PREVIOUS_REUSE_COL = 16  # P
    PREVIOUS_RESULT_COL = 17  # Q
    REMARK_COL = 15  # O

    @staticmethod
    def find_column(ws, names, default_col, search_rows=10):
        if isinstance(names, str):
            names = [names]
        for row in range(1, search_rows + 1):
            for col in range(1, ws.max_column + 1):
                value = ws.cell(row=row, column=col).value
                if value is not None and any(name in str(value) for name in names):
                    return col
        return default_col

    @staticmethod
    def _split_sheet(wb):
        exact_names = ["功能点拆分表", "2、功能点拆分表", "拆分表"]
        for name in exact_names:
            if name in wb.sheetnames:
                return wb[name]
        for name in wb.sheetnames:
            if "功能点拆分" in name or "拆分表" in name:
                return wb[name]
        raise ValueError("无法在文件中找到功能点拆分表工作表")

    @classmethod
    def _build_rows(cls, ws, func_col, desc_col):
        """Return row metadata keyed by an occurrence-aware process key.

        Excel returns None for non-top-left cells in merged functional-process cells.
        We forward-fill that process name.  A process can also contain identical H
        values, so the appearance index is part of the key; this prevents VLOOKUP
        from resolving every duplicate to the first row.
        """
        rows = {}
        current_process = ""
        occurrences = {}
        for row in range(cls.START_ROW, ws.max_row + 1):
            process_value = ws.cell(row=row, column=func_col).value
            if process_value is not None and str(process_value).strip():
                current_process = str(process_value).strip()

            description_value = ws.cell(row=row, column=desc_col).value
            if description_value is None or not str(description_value).strip():
                continue
            description = str(description_value).strip()
            base_key = (current_process, description)
            occurrences[base_key] = occurrences.get(base_key, 0) + 1
            rows[row] = {
                "key": (
                    f"{len(current_process)}:{current_process}"
                    f"{len(description)}:{description}#{occurrences[base_key]}"
                ),
                "process": current_process,
                "description": description,
            }
        return rows

    @staticmethod
    def _replace_first_number(value, replacement):
        text = str(value).strip()
        match = re.search(r"\d+", text)
        if not match:
            return text
        start, end = match.span()
        if text == match.group():
            return replacement
        return text[:start] + str(replacement) + text[end:]

    @classmethod
    def process_re_review(cls, excel1_path, excel2_path, output_dir, progress_callback=None):
        """Process Excel1 (old assessment report) and Excel2 (new re-review receipt).

        Excel2 P/Q are calculated from the old L/O columns first; then Excel1 T/U
        VLOOKUP those two new-table results.  Missing rows are represented as N/A.
        """
        if progress_callback:
            progress_callback(5)

        wb1_formula = openpyxl.load_workbook(excel1_path)
        wb1_value = openpyxl.load_workbook(excel1_path, data_only=True)
        wb2_value = openpyxl.load_workbook(excel2_path, data_only=True)
        wb2_final = openpyxl.load_workbook(excel2_path)
        try:
            if progress_callback:
                progress_callback(25)

            ws1_formula = cls._split_sheet(wb1_formula)
            ws1_value = cls._split_sheet(wb1_value)
            ws2_value = cls._split_sheet(wb2_value)
            ws2_final = cls._split_sheet(wb2_final)

            desc_col1 = cls.find_column(ws1_value, ["子过程描述", "功能点拆分描述", "子过程"], 8)
            func_col1 = cls.find_column(ws1_value, ["功能过程", "功能简述", "功能点"], 6)
            desc_col2 = cls.find_column(ws2_value, ["子过程描述", "功能点拆分描述", "子过程"], 8)
            func_col2 = cls.find_column(ws2_value, ["功能过程", "功能简述", "功能点"], 6)

            excel1_rows = cls._build_rows(ws1_value, func_col1, desc_col1)
            excel2_rows = cls._build_rows(ws2_value, func_col2, desc_col2)

            def group_key(data):
                return data["process"], data["description"]

            old_groups = {}
            new_groups = {}
            for row, data in excel1_rows.items():
                old_groups.setdefault(group_key(data), []).append(row)
            for row, data in excel2_rows.items():
                new_groups.setdefault(group_key(data), []).append(row)

            # Match each new row to the nearest unused old row in the same merged-G
            # process / H-description group.  This is deliberately not a plain
            # VLOOKUP: duplicate H values must not all resolve to the first row.
            new_to_old = {}
            assigned_old_rows = set()
            for new_row, new_data in excel2_rows.items():
                candidates = old_groups.get(group_key(new_data), [])
                available = [row for row in candidates if row not in assigned_old_rows]
                selected_from = available or candidates
                if selected_from:
                    old_row = min(selected_from, key=lambda row: (abs(row - new_row), row))
                    new_to_old[new_row] = old_row
                    assigned_old_rows.add(old_row)

            def reuse_category(value):
                text = str(value).strip() if value is not None else ""
                if "新增" in text:
                    return "新增"
                if "复用" in text:
                    return "复用"
                if "利旧" in text:
                    return "利旧"
                return "N/A"

            # New receipt P/Q are the source of truth.  For a referenced old row,
            # use the earliest new row in that same G/H group as the new identifier.
            ws2_final.cell(row=cls.START_ROW - 1, column=cls.PREVIOUS_REUSE_COL).value = "上一轮复用度"
            ws2_final.cell(row=cls.START_ROW - 1, column=cls.PREVIOUS_RESULT_COL).value = "新标号"
            new_results_by_old_row = {}
            for new_row in range(cls.START_ROW, ws2_value.max_row + 1):
                old_row = new_to_old.get(new_row)
                category = reuse_category(ws1_value.cell(row=old_row, column=12).value) if old_row else "N/A"
                if category == "N/A":
                    new_identifier = "N/A"
                elif category == "新增":
                    new_identifier = 0
                else:
                    remark = ws1_value.cell(row=old_row, column=cls.REMARK_COL).value
                    text = str(remark).strip() if remark is not None else ""
                    if category == "利旧" and not text:
                        new_identifier = "N"
                    elif not text:
                        new_identifier = "N/A"
                    else:
                        match = re.search(r"\d+", text)
                        reference_data = excel1_rows.get(int(match.group())) if match else None
                        target_rows = new_groups.get(group_key(reference_data), []) if reference_data else []
                        # Directly fill the earliest reusable row for duplicate H values.
                        new_identifier = (
                            cls._replace_first_number(text, min(target_rows))
                            if target_rows else text
                        )
                ws2_final.cell(row=new_row, column=cls.PREVIOUS_REUSE_COL).value = category
                ws2_final.cell(row=new_row, column=cls.PREVIOUS_RESULT_COL).value = new_identifier
                if old_row is not None and old_row not in new_results_by_old_row:
                    new_results_by_old_row[old_row] = (new_row, category, new_identifier)

            if progress_callback:
                progress_callback(60)

            if "Sheet1" in wb1_formula.sheetnames:
                del wb1_formula["Sheet1"]
            helper = wb1_formula.create_sheet("Sheet1")
            helper.cell(row=3, column=1).value = "旧表行号"
            helper.cell(row=3, column=2).value = "旧表唯一键"
            helper.cell(row=3, column=3).value = "新表行号"
            helper.cell(row=3, column=4).value = "新表唯一键"
            helper.cell(row=3, column=5).value = "新表上一轮复用度"
            helper.cell(row=3, column=6).value = "新表新标号"
            for old_row, data in excel1_rows.items():
                helper.cell(row=old_row, column=1).value = old_row
                helper.cell(row=old_row, column=2).value = data["key"]
                result = new_results_by_old_row.get(old_row)
                if result:
                    new_row, category, new_identifier = result
                    helper.cell(row=old_row, column=3).value = new_row
                    helper.cell(row=old_row, column=4).value = excel2_rows[new_row]["key"]
                    helper.cell(row=old_row, column=5).value = category
                    helper.cell(row=old_row, column=6).value = new_identifier

            # Excel1 T/U only VLOOKUP the two computed results from the new receipt.
            ws1_formula.cell(row=cls.START_ROW - 1, column=cls.CURRENT_REUSE_COL).value = "最新复用度"
            ws1_formula.cell(row=cls.START_ROW - 1, column=cls.NEW_ID_COL).value = "新标号"
            for old_row in range(cls.START_ROW, ws1_formula.max_row + 1):
                ws1_formula.cell(row=old_row, column=cls.CURRENT_REUSE_COL).value = (
                    f'=IFERROR(VLOOKUP(ROW(),Sheet1!$A:$F,5,FALSE),"N/A")'
                )
                ws1_formula.cell(row=old_row, column=cls.NEW_ID_COL).value = (
                    f'=IFERROR(VLOOKUP(ROW(),Sheet1!$A:$F,6,FALSE),"N/A")'
                )

            if progress_callback:
                progress_callback(85)
            os.makedirs(output_dir, exist_ok=True)
            path1 = os.path.join(output_dir, os.path.basename(excel1_path))
            path2 = os.path.join(output_dir, os.path.basename(excel2_path))
            wb1_formula.save(path1)
            wb2_final.save(path2)
            if progress_callback:
                progress_callback(100)
            return path1, path2
        finally:
            wb1_formula.close()
            wb1_value.close()
            wb2_value.close()
            wb2_final.close()
