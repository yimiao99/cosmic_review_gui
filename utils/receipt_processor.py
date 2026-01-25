import os
import re
import pandas as pd
import openpyxl
from datetime import datetime
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_COLOR_INDEX, WD_LINE_SPACING
from docx.shared import Pt
from PySide6.QtCore import QThread, Signal
from extend.matcher_config import MatcherConfig


class ReceiptWorker(QThread):
    """异步处理回单生成的线程"""

    item_finished = Signal(int, dict)  # 单个子项完成信号 (index, result)
    finished = Signal(dict)  # 改为发送字典，包含路径和统计数据
    error = Signal(str)

    def __init__(self, data, project_name=None):
        super().__init__()
        self.data = data
        self.project_name = project_name or "Receipt"

    def run(self):
        from utils.runtime_logger import RuntimeLogger

        # 设置项目名称用于日志
        RuntimeLogger.set_project(self.project_name)
        RuntimeLogger.log("开始生成回单...")

        try:
            # 检查是否为批量任务
            is_batch = self.data.get("is_batch", False)
            if is_batch and "tasks" in self.data:
                tasks = self.data["tasks"]
                last_result = None
                for i, task_data in enumerate(tasks):
                    try:
                        # 逐个生成
                        result = ReceiptProcessor.generate(task_data)
                        self.item_finished.emit(i, result)
                        last_result = result
                    except Exception as sub_e:
                        print(f"批量任务第 {i} 项生成失败: {sub_e}")
                        self.item_finished.emit(i, {"error": str(sub_e)})

                # 全部完成后发送最后的结果 (or a summary)
                self.finished.emit(last_result or {"error": "无有效任务"})
            else:
                # 单个任务
                result = ReceiptProcessor.generate(self.data)
                self.finished.emit(result)
        except Exception as e:
            import traceback

            traceback.print_exc()
            self.error.emit(str(e))


class ReceiptProcessor:
    """回单生成处理器"""

    @staticmethod
    def generate(data):
        is_merge = data.get("is_merge", False)

        if is_merge and data.get("file_groups"):
            # 合并模式:处理多个子项目
            return ReceiptProcessor._generate_merged(data)
        else:
            # 单项目模式:保持现有逻辑
            return ReceiptProcessor._generate_single(data)

    @staticmethod
    def _generate_single(data):
        """单项目模式:原有的处理逻辑"""
        # 1. 提取评估报告数据 (如果存在)
        fp_data = {}
        if data.get("eval_report_path"):
            fp_data = ReceiptProcessor._parse_eval_report(data["eval_report_path"])
        else:
            fp_data = {
                "total_fp": "N/A",
                "new_fp": "N/A",
                "reuse_fp": "N/A",
                "legacy_fp": "N/A",
                "new_ratio": "N/A",
                "reuse_ratio": "N/A",
                "legacy_ratio": "N/A",
                "total_ratio": "N/A",
            }

        # 2. 提取评估认同表数据 (如果存在)
        consent_data = {}
        if data.get("eval_consent_path"):
            consent_data = ReceiptProcessor._parse_consent_form(
                data["eval_consent_path"]
            )
        else:
            consent_data = {
                "submission_days": "0.00",
                "eval_days": "0.00",
                "reduction_ratio": "0.00%",
            }

        # 【优化】如果评估报告中有“送审人天”和“核定人天”，优先使用报告的数据
        # 即使认同表里有数据，也可能因为单元格不规范导致解析失败，报告往往更准
        if (
            fp_data.get("submission_days")
            and str(fp_data.get("submission_days")) != "0"
        ):
            consent_data["submission_days"] = fp_data["submission_days"]
        if fp_data.get("eval_days") and str(fp_data.get("eval_days")) != "0":
            consent_data["eval_days"] = fp_data["eval_days"]

        # 重新计算核减率
        try:
            sub = float(str(consent_data.get("submission_days", "0")).replace(",", ""))
            ev = float(str(consent_data.get("eval_days", "0")).replace(",", ""))
            if sub > 0:
                ratio = (sub - ev) / sub
                consent_data["reduction_ratio"] = f"{ratio:.2%}"
            else:
                consent_data["reduction_ratio"] = "0.00%"
        except:
            pass

        # 3. 选择模式 (严格对应按钮选择)
        mode = str(data.get("submission_mode", "线上"))  # 默认线上

        # 4. 合并所有需要填充的数据
        # 【新逻辑】清理项目名称中的冗余后缀（如 评估报告、日期、项目编号等）
        display_name = data["project_name"].strip()

        # 1. 递归清理常见关键词和噪音
        trash_patterns = [
            r"评估报告\s*$",
            r"结论认同表\s*$",
            r"认同表\s*$",
            r"核定表\s*$",
            r"确认单\s*$",
            r"202[4-6][-_]?\d{2,4}\s*$",  # 完整日期如 2026-0124
            r"[-_—]?\d{4,5}\s*$",  # 结尾的 4-5 位数字如 -1202 或 1202
            r"[-_]?V\d+(\.\d+)?\s*$",  # 版本号如 -V1.0
            r"副本$",
            r"\(副本\)",
        ]
        for pattern in trash_patterns:
            display_name = re.sub(
                pattern, "", display_name, flags=re.IGNORECASE
            ).strip()

        # 2. 特殊处理：如果已经以“项目”结尾，又被误伤了或者带了空格，清理它
        display_name = display_name if display_name else data["project_name"]
        if display_name.endswith("项目"):
            pass  # 保持原样

        context = {
            "project_name": display_name,
            "project_id": data["project_id"],
            "submission_unit": data["submission_unit"],
            "submitter": data["submitter"],
            "submission_time": data["submission_time"],
            "submission_mode": mode,
            "_consent_path": data["eval_consent_path"],  # 保留路径用于 Excel 公式
            **fp_data,
            **consent_data,
        }

        template_name = f"XXXXXXX项目评估确认单 - {mode}.docx"
        template_path = os.path.join("folder", template_name)
        if not os.path.exists(template_path):
            # 兼容性处理
            template_path = os.path.join(
                os.path.dirname(__file__), "..", "folder", template_name
            )

        if not os.path.exists(template_path):
            raise FileNotFoundError(f"找不到模板文件: {template_path}")

        # 5. 填充模板并保存
        doc = Document(template_path)
        ReceiptProcessor._fill_template(doc, context)

        # 【优化】文件名逻辑：使用清理后的 context["project_name"] 并防止“项目项目”
        final_display_name = context["project_name"]
        if final_display_name.endswith("项目"):
            output_name = f"{final_display_name}评估确认单.docx"
        else:
            output_name = f"{final_display_name}项目评估确认单.docx"

        # 从配置中加载回单存放位置
        config = MatcherConfig.load()
        output_dir = config.get("storage", {}).get("receipt")

        if not output_dir:
            # 兼容：优先保存在报告所在目录
            output_dir = (
                os.path.dirname(data["eval_report_path"])
                if data.get("eval_report_path")
                else "."
            )

        final_path = os.path.join(output_dir, output_name)
        counter = 1
        while True:
            try:
                # 检查输出目录是否存在
                if not os.path.exists(output_dir):
                    os.makedirs(output_dir)
                doc.save(final_path)
                break
            except (IOError, PermissionError):
                # 如果文件正在被打开，则生成副本
                base, ext = os.path.splitext(output_name)
                final_path = os.path.join(output_dir, f"{base}({counter}){ext}")
                counter += 1

        # 6. 回写结果到评估报告 PQR 列 (使用副本回写以避免被占用)
        if data.get("eval_report_path"):
            try:
                ReceiptProcessor._write_back_to_excel(data["eval_report_path"], context)
            except RuntimeError as e:
                # 如果回写失败（通常是文件占用），生成一个带 _结果 的副本
                base, ext = os.path.splitext(data["eval_report_path"])
                copy_path = f"{base}_回写结果{ext}"
                import shutil

                try:
                    shutil.copy2(data["eval_report_path"], copy_path)
                    ReceiptProcessor._write_back_to_excel(copy_path, context)
                    # 如果回写成功，通知或者记录日志（由于目前是直接返回 output_path，这里仅作为防御逻辑）
                except Exception:
                    pass

        # 7. 计算统计数据用于卡片显示
        stats = {
            "new_fp": fp_data.get("new_fp", "N/A"),
            "reuse_fp": fp_data.get("reuse_fp", "N/A"),
            "legacy_fp": fp_data.get("legacy_fp", "N/A"),
            "total_fp": fp_data.get("total_fp", "N/A"),
            "new_ratio": fp_data.get("new_ratio", "0.0%"),
            "reuse_ratio": fp_data.get("reuse_ratio", "0.0%"),
            "legacy_ratio": fp_data.get("legacy_ratio", "0.0%"),
            "total_ratio": fp_data.get("total_ratio", "100.0%"),
            "submission_days": consent_data.get("submission_days", "0.00"),
            "eval_days": consent_data.get("eval_days", "0.00"),
            "reduction_ratio": consent_data.get("reduction_ratio", "0.00%"),
        }

        return {
            "output_path": final_path,
            "stats": stats,
            "excel_reports": (
                [data.get("eval_report_path")] if data.get("eval_report_path") else []
            ),
        }

    @staticmethod
    def _extract_sub_project_name(consent_path, project_name):
        """
        提取子项目名称
        优先级:
        1. 从结论认同表A3单元格提取
        2. 从文件名提取(去掉项目名称前缀)
        """
        sub_name = ""
        try:
            wb = openpyxl.load_workbook(consent_path, data_only=True)
            # 尝试从第一个sheet的A3获取
            ws = wb.active
            a3_value = ws["A3"].value
            wb.close()

            if a3_value and isinstance(a3_value, str) and a3_value.strip():
                # 去掉项目名称前缀
                candidate = a3_value.strip()
                # 过滤掉明显的无效名称
                invalid_keywords = [
                    "功能点拆分",
                    "功能规模",
                    "确认单",
                    "认同表",
                    "计算公式",
                ]
                if not any(k in candidate for k in invalid_keywords):
                    sub_name = candidate
                    if project_name in sub_name:
                        sub_name = sub_name.replace(project_name, "").strip()
        except:
            pass

        if not sub_name:
            # 备选:从文件名提取
            filename = os.path.basename(consent_path)
            # 去掉扩展名和项目名称
            sub_name = filename.replace(".xlsx", "").replace(".xls", "")
            if project_name in sub_name:
                sub_name = sub_name.replace(project_name, "").strip()

            # 进一步清理文件名中的垃圾信息
            trash_words = ["结论认同表", "核定单", "评估报告", "V1", "V2", "副本"]
            for tw in trash_words:
                sub_name = sub_name.replace(tw, "")

        # 【优化】清理常见的连接符和括号
        if sub_name:
            sub_name = re.sub(r"^[-_—\s]+", "", sub_name)
            sub_name = re.sub(r"[-_—\s]+$", "", sub_name)
            if (sub_name.startswith("(") and sub_name.endswith(")")) or (
                sub_name.startswith("（") and sub_name.endswith("）")
            ):
                sub_name = sub_name[1:-1].strip()

        return sub_name or "子项目"

    @staticmethod
    def _generate_merged(data):
        """
        处理合并模式:多个子项目
        """
        from utils.runtime_logger import RuntimeLogger

        file_groups = data["file_groups"]
        # 【新逻辑】清理项目名称
        project_name = data["project_name"].strip()
        trash_patterns = [
            r"评估报告\s*$",
            r"结论认同表\s*$",
            r"认同表\s*$",
            r"核定表\s*$",
            r"确认单\s*$",
            r"202[4-6][-_]?\d{2,4}\s*$",
            r"[-_—]?\d{4,5}\s*$",
            r"[-_]?V\d+(\.\d+)?\s*$",
            r"副本$",
        ]
        for pattern in trash_patterns:
            project_name = re.sub(
                pattern, "", project_name, flags=re.IGNORECASE
            ).strip()

        project_name = project_name if project_name else data["project_name"]
        project_name = project_name.strip(" -_—")

        RuntimeLogger.log(
            f"开始处理合并任务: {project_name}, 子项目数量: {len(file_groups)}"
        )

        # 调试日志：列出文件组
        for gk, fv in file_groups.items():
            RuntimeLogger.log(
                f"待处理文件组: {gk} -> {os.path.basename(fv.get('eval_report',''))}"
            )

        sub_projects = []
        all_stats = []

        for group_name, files in file_groups.items():
            report_path = files.get("eval_report")
            consent_path = files.get("eval_consent")

            if not report_path or not consent_path:
                RuntimeLogger.log(f"跳过不完整的子项目组: {group_name}", "warning")
                continue

            # 1. 提取子项目名称
            sub_name = ReceiptProcessor._extract_sub_project_name(
                consent_path, project_name
            )
            RuntimeLogger.log(f"--- 处理子项目: {sub_name} ---")

            # 2. 解析数据
            fp_data = ReceiptProcessor._parse_eval_report(report_path)
            consent_data = ReceiptProcessor._parse_consent_form(consent_path)

            # 3. 数据融合：优先使用报告里的人天数据
            if fp_data.get("submission_days"):
                val = str(fp_data.get("submission_days")).strip()
                if (
                    val
                    and val.lower() != "none"
                    and val.lower() != "n/a"
                    and val != "0"
                    and val != "0.00"
                ):
                    consent_data["submission_days"] = val
                    RuntimeLogger.log(f"已采用评估报告中的送审人天: {val}")

            if fp_data.get("eval_days"):
                val = str(fp_data.get("eval_days")).strip()
                if (
                    val
                    and val.lower() != "none"
                    and val.lower() != "n/a"
                    and val != "0"
                    and val != "0.00"
                ):
                    consent_data["eval_days"] = val
                    RuntimeLogger.log(f"已采用评估报告中的核定人天: {val}")

            # 4. 计算核减率 (针对单个子项目)
            try:
                sub_val = float(
                    str(consent_data.get("submission_days", 0))
                    .replace(",", "")
                    .replace("人天", "")
                )
                ev_val = float(
                    str(consent_data.get("eval_days", 0))
                    .replace(",", "")
                    .replace("人天", "")
                )
                if sub_val > 0:
                    reduction_ratio = f"{(sub_val - ev_val) / sub_val:.2%}"
                else:
                    reduction_ratio = "0.00%"
            except:
                reduction_ratio = "0.00%"

            # 5. 组装子项目数据 (用于 Word)
            sub_project_data = {
                "sub_project_name": sub_name,
                "submission_fp": fp_data.get("total_fp", 0),
                "submission_days": consent_data.get("submission_days", "0"),
                "new_fp": fp_data.get("new_fp", 0),
                "new_ratio": fp_data.get("new_ratio", "0%"),
                "reuse_fp": fp_data.get("reuse_fp", 0),
                "reuse_ratio": fp_data.get("reuse_ratio", "0%"),
                "legacy_fp": fp_data.get("legacy_fp", 0),
                "legacy_ratio": fp_data.get("legacy_ratio", "0%"),
                "eval_days": consent_data.get("eval_days", "0"),
            }
            sub_projects.append(sub_project_data)

            # 6. 组装统计数据 (用于 UI)
            stats = {
                "sub_project_name": sub_name,
                "new_fp": fp_data.get("new_fp", 0),
                "reuse_fp": fp_data.get("reuse_fp", 0),
                "legacy_fp": fp_data.get("legacy_fp", 0),
                "total_fp": fp_data.get("total_fp", 0),
                "new_ratio": fp_data.get("new_ratio", "0.0%"),
                "reuse_ratio": fp_data.get("reuse_ratio", "0.0%"),
                "legacy_ratio": fp_data.get("legacy_ratio", "0.0%"),
                "submission_days": consent_data.get("submission_days", "0.00"),
                "eval_days": consent_data.get("eval_days", "0.00"),
                "reduction_ratio": reduction_ratio,
            }
            all_stats.append(stats)

            # 7. 回写结果到 Excel
            context_to_write = {
                "project_name": project_name,
                "_consent_path": consent_path,
                **fp_data,
                **consent_data,
                "reduction_ratio": reduction_ratio,
            }
            try:
                ReceiptProcessor._write_back_to_excel(report_path, context_to_write)
            except Exception as e:
                RuntimeLogger.log(f"回写Excel失败: {str(e)}", "warning")

        # 8. 计算汇总数据
        total_stats = ReceiptProcessor._calculate_total_stats(sub_projects)
        RuntimeLogger.log(
            f"汇总计算完成: 总送审 {total_stats['total_submission_days']}, 总核定 {total_stats['total_eval_days']}, 整体核减率 {total_stats['total_reduction_ratio']}"
        )

        # 9. 将汇总数据加入 all_stats 供 UI 显示
        all_stats.append(
            {
                "is_total": True,
                "sub_project_name": "汇总合计",
                "new_fp": total_stats["new_fp"],
                "reuse_fp": total_stats["reuse_fp"],
                "legacy_fp": total_stats["legacy_fp"],
                "total_fp": total_stats["total_fp"],
                "new_ratio": total_stats["new_ratio"],
                "reuse_ratio": total_stats["reuse_ratio"],
                "legacy_ratio": total_stats["legacy_ratio"],
                "submission_days": total_stats["total_submission_days"],
                "eval_days": total_stats["total_eval_days"],
                "reduction_ratio": total_stats["total_reduction_ratio"],
            }
        )

        # 10. 填充 Word 模板
        context_word = {
            "project_name": project_name,
            "project_id": data["project_id"],
            "submission_unit": data["submission_unit"],
            "submitter": data["submitter"],
            "submission_time": data["submission_time"],
            "submission_mode": "合并",
            "is_merge": True,
            "sub_projects": sub_projects,
            "sub_project_count": len(sub_projects),
            "submission_days": total_stats["total_submission_days"],
            "eval_days": total_stats["total_eval_days"],
            "reduction_ratio": total_stats["total_reduction_ratio"],
            **total_stats,
        }

        template_name = "XXXXXXX项目评估确认单 - 合并.docx"
        template_path = os.path.join("folder", template_name)
        if not os.path.exists(template_path):
            template_path = os.path.join(
                os.path.dirname(__file__), "..", "folder", template_name
            )

        if not os.path.exists(template_path):
            RuntimeLogger.log(f"找不到合并模板: {template_path}", "error")
            raise FileNotFoundError(f"找不到合并模板: {template_path}")

        doc = Document(template_path)
        ReceiptProcessor._fill_template(doc, context_word)

        # 11. 保存 Word
        config = MatcherConfig.load()
        output_dir = config.get("storage", {}).get("receipt", ".")

        # 【优化】文件名逻辑：防止“项目项目”
        if project_name.endswith("项目"):
            output_name = f"{project_name}评估确认单.docx"
        else:
            output_name = f"{project_name}项目评估确认单.docx"

        final_path = os.path.join(output_dir, output_name)

        counter = 1
        while True:
            try:
                if not os.path.exists(output_dir):
                    os.makedirs(output_dir)
                doc.save(final_path)
                break
            except (IOError, PermissionError):
                base, ext = os.path.splitext(output_name)
                final_path = os.path.join(output_dir, f"{base}({counter}){ext}")
                counter += 1

        RuntimeLogger.log(f"确认单生成成功: {os.path.basename(final_path)}")

        excel_reports = [f["eval_report"] for f in file_groups.values()]
        return {
            "output_path": final_path,
            "stats": all_stats,
            "excel_reports": excel_reports,
        }

    @staticmethod
    def _calculate_total_stats(sub_projects):
        """
        计算所有子项目的汇总数据
        """
        total_submission_days = 0.0
        total_eval_days = 0.0
        total_new_fp = 0
        total_reuse_fp = 0
        total_legacy_fp = 0

        for sub in sub_projects:
            try:
                days = (
                    str(sub.get("submission_days", "0"))
                    .replace(",", "")
                    .replace("人天", "")
                )
                total_submission_days += float(days)
            except:
                pass
            try:
                days = (
                    str(sub.get("eval_days", "0")).replace(",", "").replace("人天", "")
                )
                total_eval_days += float(days)
            except:
                pass
            try:
                total_new_fp += int(float(str(sub["new_fp"]).replace(",", "")))
            except:
                pass
            try:
                total_reuse_fp += int(float(str(sub["reuse_fp"]).replace(",", "")))
            except:
                pass
            try:
                total_legacy_fp += int(float(str(sub["legacy_fp"]).replace(",", "")))
            except:
                pass

        total_fp = total_new_fp + total_reuse_fp + total_legacy_fp

        def safe_ratio(part, total):
            if total == 0:
                return "0.0%"
            return f"{(part / total):.1%}"

        # 计算核减比例
        if total_submission_days > 0:
            reduction_ratio = (1 - total_eval_days / total_submission_days) * 100
        else:
            reduction_ratio = 0.0

        return {
            "total_submission_days": f"{total_submission_days:.2f}",
            "total_eval_days": f"{total_eval_days:.2f}",
            "total_reduction_ratio": f"{reduction_ratio:.2f}%",
            "new_fp": total_new_fp,
            "reuse_fp": total_reuse_fp,
            "legacy_fp": total_legacy_fp,
            "total_fp": total_fp,
            "new_ratio": safe_ratio(total_new_fp, total_fp),
            "reuse_ratio": safe_ratio(total_reuse_fp, total_fp),
            "legacy_ratio": safe_ratio(total_legacy_fp, total_fp),
        }

    @staticmethod
    def _write_back_to_excel(path, context):
        """将统计结果和结论表数据一并回写到评估报告 Excel 的 PQR 列"""
        try:
            wb = openpyxl.load_workbook(path)
            # 查找功能点拆分表
            ws = None
            for sheet in wb.sheetnames:
                if "功能点拆分表" in sheet:
                    ws = wb[sheet]
                    break
            if not ws:
                ws = wb.active

            # 找到起始行 (从 Q 列第 10 行开始寻找第一个空行)
            start_row = 10
            for r in range(10, 100):
                if (
                    ws.cell(row=r, column=17).value is None
                    and ws.cell(row=r + 1, column=17).value is None
                ):
                    start_row = r
                    break

            # P(16): 标签, Q(17): 数值(公式), R(18): 占比(公式)
            col_letter = openpyxl.utils.get_column_letter(12)  # L
            total_r_idx = start_row + 3

            # 1. 功能点公式回写
            fp_rows = [
                (
                    "新增",
                    f'=COUNTIF({col_letter}:{col_letter}, "新增")',
                    f"=Q{{row}}/Q{total_r_idx}",
                ),
                (
                    "复用",
                    f'=COUNTIF({col_letter}:{col_letter}, "复用")',
                    f"=Q{{row}}/Q{total_r_idx}",
                ),
                (
                    "利旧",
                    f'=COUNTIF({col_letter}:{col_letter}, "利旧")',
                    f"=Q{{row}}/Q{total_r_idx}",
                ),
                ("合计", f"=Q{start_row}+Q{start_row+1}+Q{start_row+2}", "100.0%"),
            ]

            curr_r = start_row
            for i, (label, val_fmt, ratio_fmt) in enumerate(fp_rows):
                ws.cell(row=curr_r, column=16).value = label
                ws.cell(row=curr_r, column=17).value = val_fmt
                if i < 3:
                    ws.cell(row=curr_r, column=18).value = ratio_fmt.replace(
                        "{row}", str(curr_r)
                    )
                else:
                    ws.cell(row=curr_r, column=18).value = "100.0%"
                curr_r += 1

            # 2. 结论认同表数值回写 (使用绝对路径公式，解决闭卷 #REF! 问题)
            curr_r += 2
            consent_path = context.get("_consent_path")
            if consent_path:
                abs_path = os.path.abspath(consent_path)
                dir_name = os.path.dirname(abs_path)
                file_name = os.path.basename(abs_path)
                sheet_name = context.get("found_sheet") or "结论认同表"

                # Excel 闭卷引用格式: ='C:\路径\[文件名.xlsx]工作表'!$A$1
                # 注意：跨文件引用必须使用绝对路径，否则关闭源文件后会显示 #REF!
                ref_prefix = f"='{dir_name}\\[{file_name}]{sheet_name}'"

                consent_data = [
                    (
                        "送审人天",
                        f"{ref_prefix}!{context.get('submission_addr', '$C$3')}",
                        "",
                    ),
                    (
                        "核定人天",
                        f"{ref_prefix}!{context.get('eval_addr', '$D$3')}",
                        "",
                    ),
                    (
                        "核减比例",
                        f"{ref_prefix}!{context.get('ratio_addr', '$E$3')}",
                        "",
                    ),
                ]
            else:
                consent_data = [
                    ("送审人天", context.get("submission_days", "0"), ""),
                    ("核定人天", context.get("eval_days", "0.00"), ""),
                    ("核减比例", context.get("reduction_ratio", "0.0%"), ""),
                ]

            for label, val, ratio in consent_data:
                ws.cell(row=curr_r, column=16).value = label
                ws.cell(row=curr_r, column=17).value = val
                ws.cell(row=curr_r, column=18).value = ratio
                curr_r += 1

            wb.save(path)
        except PermissionError:
            raise RuntimeError(
                f"无法保存 Excel：文件已被打开，请关闭 {os.path.basename(path)}"
            )
        except Exception as e:
            raise RuntimeError(f"Excel 回写失败: {str(e)}")

    @staticmethod
    def _parse_eval_report(path):
        """解析评估报告获取项目统计 (严格基于复用度列行数统计)"""
        from utils.runtime_logger import RuntimeLogger

        RuntimeLogger.log(f"解析评估报告: {os.path.basename(path)}")

        try:
            df_dict = pd.read_excel(path, sheet_name=None)
        except Exception as e:
            RuntimeLogger.log(f"无法读取评估报告: {str(e)}", "error")
            raise RuntimeError(f"无法读取评估报告: {str(e)}")

        # 尝试查找所有可能的数据点
        report_results = {
            "total_fp": 0,
            "new_fp": 0,
            "reuse_fp": 0,
            "legacy_fp": 0,
            "new_ratio": "0.0%",
            "reuse_ratio": "0.0%",
            "legacy_ratio": "0.0%",
            "submission_days": None,
            "eval_days": None,
        }

        # 1. 首先尝试从“汇总”或带有关键词的工作表提取人天
        found_sub = False
        found_eval = False

        all_sheets = list(df_dict.keys())
        RuntimeLogger.log(f"开始提取人天数据，所有工作表: {all_sheets}")

        for name, df in df_dict.items():
            if found_sub and found_eval:
                break

            # 宽松匹配 Sheet 名
            is_summary_sheet = any(
                k in name
                for k in [
                    "汇总",
                    "合计",
                    "结果",
                    "评估",
                    "认同",
                    "统计",
                    "分值",
                    "Sheet1",
                    "录入",
                ]
            )
            if is_summary_sheet:
                RuntimeLogger.log(f"正在扫描潜力工作表: {name} ...")
                # 遍历前100行找“送审人天”等关键字
                for r in range(min(100, len(df))):
                    if found_sub and found_eval:
                        break

                    row_vals = [str(x) for x in df.iloc[r].values]
                    for idx, val in enumerate(row_vals):
                        if not val or val == "nan":
                            continue
                        clean_v = (
                            val.replace(" ", "").replace("\n", "").replace("\r", "")
                        )

                        target_key = None
                        if (
                            any(
                                k in clean_v
                                for k in [
                                    "送审人天",
                                    "送审工作量",
                                    "送审自评",
                                    "项目送审",
                                ]
                            )
                            and not found_sub
                        ):
                            target_key = "submission_days"
                        elif (
                            any(
                                k in clean_v
                                for k in [
                                    "核定人天",
                                    "核定工作量",
                                    "评估工作量",
                                    "评估人天",
                                    "评定工作量",
                                    "评定结果",
                                    "评估结果",
                                ]
                            )
                            and not found_eval
                        ):
                            target_key = "eval_days"

                        if target_key:
                            # 尝试从右侧提取
                            val_found = None
                            try:
                                # 扫描右侧 6 个单元格
                                for offset in range(1, 7):
                                    if idx + offset >= len(df.columns):
                                        break
                                    candidate = str(df.iloc[r, idx + offset]).strip()
                                    if not candidate or candidate == "nan":
                                        continue
                                    # 提取数字 (允许带千分位)
                                    clean_cand = candidate.replace(",", "")
                                    num_match = re.search(r"(\d+(\.\d+)?)", clean_cand)
                                    if num_match:
                                        cur_val = float(num_match.group(1))
                                        if cur_val > 0:
                                            val_found = f"{cur_val:.2f}"
                                            break
                            except:
                                pass

                            # 如果右侧没找到有力数值，看看当前单元格
                            if not val_found:
                                # 尝试从当前单元格正则提取 (如 "送审工作量：123.45")
                                m = re.search(
                                    r"[:：]\s*(\d+(?:\.\d+)?)", clean_v
                                ) or re.search(r"(\d+(?:\.\d+)?)", clean_v)
                                if m and float(m.group(1)) > 0:
                                    val_found = f"{float(m.group(1)):.2f}"

                            if val_found:
                                report_results[target_key] = val_found
                                if target_key == "submission_days":
                                    found_sub = True
                                else:
                                    found_eval = True
                                RuntimeLogger.log(
                                    f"-> 在 [{name}] 表中找到 {target_key}: {val_found}"
                                )
                                break

        # 1.5 如果汇总没找齐，尝试暴力扫描所有 sheet！
        if not found_sub or not found_eval:
            RuntimeLogger.log(
                "汇总表中未找齐数值，开始全局扫描所有工作表...", "warning"
            )
            for name, df in df_dict.items():
                if found_sub and found_eval:
                    break
                for r in range(min(100, len(df))):
                    if found_sub and found_eval:
                        break
                    row_vals = [str(x) for x in df.iloc[r].values]
                    for idx, val in enumerate(row_vals):
                        clean_v = (
                            val.replace(" ", "").replace("\n", "").replace("\r", "")
                        )
                        tk = None
                        if not found_sub and any(
                            k in clean_v for k in ["送审人天", "送审工作量"]
                        ):
                            tk = "submission_days"
                        elif not found_eval and any(
                            k in clean_v
                            for k in ["核定人天", "评估工作量", "评定工作量"]
                        ):
                            tk = "eval_days"

                        if tk:
                            for offset in range(1, 4):
                                if idx + offset >= len(df.columns):
                                    break
                                cand = str(df.iloc[r, idx + offset])
                                m = re.search(r"(\d+(\.\d+)?)", cand)
                                if m and float(m.group(1)) > 0:
                                    report_results[tk] = f"{float(m.group(1)):.2f}"
                                    if tk == "submission_days":
                                        found_sub = True
                                    else:
                                        found_eval = True
                                    break

        # 2. 寻找功能点拆分表进行明细统计
        target_df = None
        for name, df in df_dict.items():
            if any(k in name for k in ["功能点", "拆分", "清单", "明细"]):
                # 排除掉太小的 sheet
                if len(df.columns) > 5 and len(df) > 5:
                    target_df = df
                    RuntimeLogger.log(f"锁定功能点数据表: {name}")
                    break

        if target_df is None:
            # 兜底：查找列数较多的表，或者是任何不叫“认同”或“结论”的表
            for name, df in df_dict.items():
                if "结论" in name or "认同" in name:
                    continue
                target_df = df
                RuntimeLogger.log(f"未能确定拆分表，尝试从 [{name}] sheet 提取")
                break

        if target_df is not None and not target_df.empty:
            # 查找“复用度”列
            type_col = -1
            start_data_idx = 0

            # 安全检查：确保 target_df 确实可以进行 len() 操作且不为空
            try:
                max_rows = min(100, len(target_df))  # 扩大扫描范围
            except:
                max_rows = 0

            for i in range(max_rows):
                row = target_df.iloc[i].values
                for idx, val in enumerate(row):
                    if isinstance(val, str) and any(
                        k in val for k in ["复用度", "开发类型", "类型", "模式"]
                    ):
                        # 进一步确认这一列下方是否包含“新增”等关键字
                        found_kw = False
                        try:
                            # 确保不越界
                            scan_end = min(i + 10, len(target_df))
                            for next_r in range(i + 1, scan_end):
                                cell_v = str(target_df.iloc[next_r, idx])
                                if any(k in cell_v for k in ["新增", "复用", "利旧"]):
                                    found_kw = True
                                    break
                        except:
                            pass

                        if found_kw:
                            type_col = idx
                            start_data_idx = i + 1
                            break
                if type_col != -1:
                    break

            if type_col != -1:
                new_count = 0
                reuse_count = 0
                legacy_count = 0

                # 统计各类型数量
                for idx in range(start_data_idx, len(target_df)):
                    row = target_df.iloc[idx]
                    val = str(row.iloc[type_col]) if type_col < len(row) else ""
                    clean_val = val.strip()
                    if (
                        not clean_val
                        or "nan" in clean_val.lower()
                        or "合计" in clean_val
                    ):
                        continue

                    if any(k in clean_val for k in ["新增", "增加"]):
                        new_count += 1
                    elif any(k in clean_val for k in ["复用", "优化", "修改", "变更"]):
                        reuse_count += 1
                    elif any(k in clean_val for k in ["利旧", "原有", "保留"]):
                        legacy_count += 1

                total_fp = new_count + reuse_count + legacy_count

                # 【优化】如果明细统计还是 0，尝试查找是否有“功能点数”列直接求和
                if total_fp == 0:
                    fp_val_col = -1
                    for idx, val in enumerate(
                        target_df.iloc[start_data_idx - 1].values
                    ):
                        if isinstance(val, str) and (
                            "功能点数" in val or "FP" in val.upper()
                        ):
                            fp_val_col = idx
                            break
                    if fp_val_col != -1:
                        try:
                            # 简单求和
                            col_vals = pd.to_numeric(
                                target_df.iloc[start_data_idx:, fp_val_col],
                                errors="coerce",
                            ).fillna(0)
                            total_fp = int(col_vals.sum())
                            # 既然分不出类型，就全部归入“新增”或者按 0 处理
                            # 这里保持 0 报警可能更好，或者至少 log 一下
                            RuntimeLogger.log(
                                f"明细复用度统计失败，但从功能点数列提取到总计: {total_fp}",
                                "warning",
                            )
                        except:
                            pass

                if total_fp > 0:
                    report_results.update(
                        {
                            "total_fp": total_fp,
                            "new_fp": new_count,
                            "reuse_fp": reuse_count,
                            "legacy_fp": legacy_count,
                            "new_ratio": (
                                f"{(new_count / total_fp):.1%}"
                                if total_fp > 0
                                else "0.0%"
                            ),
                            "reuse_ratio": (
                                f"{(reuse_count / total_fp):.1%}"
                                if total_fp > 0
                                else "0.0%"
                            ),
                            "legacy_ratio": (
                                f"{(legacy_count / total_fp):.1%}"
                                if total_fp > 0
                                else "0.0%"
                            ),
                        }
                    )
                    RuntimeLogger.log(
                        f"解析到功能点: 总计 {total_fp} (新增 {new_count}, 复用 {reuse_count}, 利旧 {legacy_count})"
                    )
            else:
                RuntimeLogger.log("未在拆分表中找到'复用度'或'开发类型'列", "warning")

        # 3. 兜底逻辑：如果全表统计下来发现功能点依然是 0，则在所有 sheet 中暴力搜索含有 "功能点" 关键字的数值单元格
        if report_results.get("total_fp", 0) == 0:
            RuntimeLogger.log(
                "常规统计 FP 失败，开始全表暴力搜索功能点总数...", "warning"
            )
            found_violence = False
            for name, df in df_dict.items():
                if found_violence:
                    break
                # 遍寻前 100 行
                for r_idx in range(min(100, len(df))):
                    if found_violence:
                        break
                    row = df.iloc[r_idx]
                    for c_idx, val in enumerate(row):
                        val_str = str(val).replace(" ", "")
                        if "功能点" in val_str and any(
                            k in val_str for k in ["总计", "合计", "数", "FP"]
                        ):
                            # 搜索该单元格右侧及下方
                            for off_r in range(0, 2):
                                if r_idx + off_r >= len(df):
                                    break
                                for off_c in range(1, 4):
                                    if c_idx + off_c >= len(df.columns):
                                        break
                                    cand = str(
                                        df.iloc[r_idx + off_r, c_idx + off_c]
                                    ).replace(",", "")
                                    try:
                                        v = float(
                                            re.search(r"(\d+(\.\d+)?)", cand).group(1)
                                        )
                                        if v > 1:  # 排除掉太小的干扰项
                                            report_results["total_fp"] = int(v)
                                            report_results["new_fp"] = int(v)
                                            report_results["new_ratio"] = "100.0%"
                                            RuntimeLogger.log(
                                                f"-> 暴力搜索在 [{name}] 表 {r_idx+off_r+1}行{c_idx+off_c+1}列 锁定 FP: {v}"
                                            )
                                            found_violence = True
                                            break
                                    except:
                                        pass
                                if found_violence:
                                    break
                            if found_violence:
                                break

        return report_results

    @staticmethod
    def _parse_consent_form(path):
        """解析结论认同表并记录数值所在的单元格地址，以便生成公式"""
        from utils.runtime_logger import RuntimeLogger

        RuntimeLogger.log(f"解析结论认同表: {os.path.basename(path)}")

        results = {
            "submission_days": "0",
            "submission_addr": "$C$3",
            "eval_days": "0.00",
            "eval_addr": "$D$3",
            "reduction_ratio": "0.0%",
            "ratio_addr": "$E$3",
            "found_sheet": None,
        }

        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            header_map = {
                "送审人天": ("submission_days", "submission_addr"),
                "核定人天": ("eval_days", "eval_addr"),
                "评估人天": ("eval_days", "eval_addr"),
                "核定工作量": ("eval_days", "eval_addr"),
                "评估工作量": ("eval_days", "eval_addr"),
                "核减比例": ("reduction_ratio", "ratio_addr"),
                "核减率": ("reduction_ratio", "ratio_addr"),
            }
            all_header_keys = list(header_map.keys())

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for r in range(1, 101):
                    for c in range(1, 31):
                        cell_val = ws.cell(row=r, column=c).value
                        if cell_val and isinstance(cell_val, str):
                            clean_text = (
                                cell_val.replace(" ", "")
                                .replace("\n", "")
                                .replace("\r", "")
                            )
                            for key, (attr, addr_attr) in header_map.items():
                                if key in clean_text:
                                    # 如果已经在一个 Sheet 找到了数据，且当前 Sheet 不是那个 Sheet，则跳过（除非还没找到 Sheet）
                                    if (
                                        results["found_sheet"]
                                        and results["found_sheet"] != sheet_name
                                    ):
                                        # 如果这个属性已经有值了，就不再从别的 Sheet 覆盖
                                        if (
                                            results[attr] != "0"
                                            and results[attr] != "0.00"
                                            and results[attr] != "0.0%"
                                        ):
                                            continue

                                    # 智能探测：全方位探测数值 (先右后下)
                                    # 1. 探测右侧单元格 (需排除右侧也是另一个表头的情况)
                                    val_right = ws.cell(row=r, column=c + 1).value
                                    is_right_valid = val_right is not None
                                    if is_right_valid and isinstance(val_right, str):
                                        rt_clean = val_right.replace(" ", "")
                                        # 如果右侧单元格内容包含任一已知表头关键字，则认为它是下一个表头
                                        if any(k in rt_clean for k in all_header_keys):
                                            is_right_valid = False

                                    # 2. 探测下方单元格 (跳过可能的表头)
                                    target_r = r + 1
                                    if target_r <= 2:
                                        target_r = 3
                                    val_below = ws.cell(row=target_r, column=c).value

                                    # 优先级：右侧(非表头) > 下方
                                    if is_right_valid and (
                                        isinstance(val_right, (int, float))
                                        or (
                                            isinstance(val_right, str)
                                            and len(val_right.strip()) > 0
                                        )
                                    ):
                                        val_cell = ws.cell(row=r, column=c + 1)
                                        target_col = c + 1
                                        target_row = r
                                    else:
                                        val_cell = ws.cell(row=target_r, column=c)
                                        target_col = c
                                        target_row = target_r

                                    results["found_sheet"] = sheet_name

                                    # 记录绝对地址，如 $C$3
                                    col_letter = openpyxl.utils.get_column_letter(
                                        target_col
                                    )
                                    results[addr_attr] = f"${col_letter}${target_row}"

                                    val = val_cell.value
                                    if val is not None:
                                        if isinstance(val, (int, float)):
                                            if "比例" in key:
                                                results[attr] = (
                                                    f"{val:.2%}"
                                                    if val < 1
                                                    else f"{val}%"
                                                )
                                            else:
                                                results[attr] = f"{float(val):.2f}"
                                        else:
                                            # 处理文本形式的百分比
                                            results[attr] = str(val).strip()

                                    RuntimeLogger.log(
                                        f"找到 {key} -> {results[attr]} (位置: {results[addr_attr]})"
                                    )
                                    break  # 找到一个关键词就跳出关键词循环
            wb.close()
        except Exception as e:
            RuntimeLogger.log(f"解析结论认同表出错: {str(e)}", "error")
        return results

    @staticmethod
    def _disable_snap_to_grid(p):
        """取消段落对齐到文档网格，解决微软雅黑中文间距过大的问题"""
        try:
            pPr = p._element.get_or_add_pPr()
            snap = pPr.find(qn("w:snapToGrid"))
            if snap is None:
                snap = OxmlElement("w:snapToGrid")
                pPr.append(snap)
            snap.set(qn("w:val"), "0")
        except Exception:
            pass

    @staticmethod
    def _fill_template(doc, context):
        """填充 Word 模板 (根据语义隔离精准匹配数据)"""

        # 0. 强制使用微软雅黑 11 号字体
        sample_font_name = "微软雅黑"
        sample_font_size = Pt(11)  # 11号字体

        # 1. 准备数据包
        def format_days(val):
            try:
                f_val = float(str(val).replace(",", "").replace("人天", ""))
                # 严格保留 2 位小数，不使用千分位
                return f"{f_val:.2f}"
            except:
                return str(val)

        def format_fp(val):
            try:
                # 功能点不使用千分位
                f_val = float(str(val).replace(",", ""))
                if f_val == int(f_val):
                    return str(int(f_val))
                return f"{f_val:.1f}"
            except:
                return str(val)

        data_bundle = {
            "total_fp": format_fp(context.get("total_fp", 0)),
            "new_fp": format_fp(context.get("new_fp", "0")),
            "new_ratio": str(context.get("new_ratio", "0%")),
            "reuse_fp": format_fp(context.get("reuse_fp", "0")),
            "reuse_ratio": str(context.get("reuse_ratio", "0%")),
            "legacy_fp": format_fp(context.get("legacy_fp", "0")),
            "legacy_ratio": str(context.get("legacy_ratio", "0%")),
            "submission_days": format_days(context.get("submission_days", "0")),
            "eval_days": format_days(context.get("eval_days", "0")),
            "reduction_ratio": str(context.get("reduction_ratio", "0%")),
        }

        if context.get("is_merge"):
            data_bundle.update(
                {
                    "total_submission_days": format_days(
                        context.get("total_submission_days", "0")
                    ),
                    "total_eval_days": format_days(context.get("total_eval_days", "0")),
                    "total_reduction_ratio": str(
                        context.get("total_reduction_ratio", "0%")
                    ),
                }
            )

        # 2. 定位并处理 Section II (评估过程)
        def handle_section_ii():
            from utils.runtime_logger import RuntimeLogger

            RuntimeLogger.log("[DEBUG] 开始执行 handle_section_ii()")

            # 打印前 30 个段落的内容用于诊断
            RuntimeLogger.log(f"[DEBUG] 文档总共有 {len(doc.paragraphs)} 个段落")
            for i in range(min(30, len(doc.paragraphs))):
                p_text = doc.paragraphs[i].text.strip()
                if p_text:  # 只打印非空段落
                    RuntimeLogger.log(
                        f"[DEBUG] 段落[{i}]: '{p_text[:50]}'"
                    )  # 只打印前50个字符

            sec_ii_para = None
            sec_iii_para = None

            for i, p in enumerate(doc.paragraphs):
                text = p.text.replace(" ", "")
                # 兼容有无编号的两种格式
                if "评估过程" in text and ("二、" in text or text == "评估过程"):
                    sec_ii_para = i
                    RuntimeLogger.log(f"[DEBUG] 找到'评估过程'在段落 {i}")
                elif "审核结果" in text and ("三、" in text or text == "审核结果"):
                    sec_iii_para = i
                    RuntimeLogger.log(f"[DEBUG] 找到'审核结果'在段落 {i}")
                    break

            RuntimeLogger.log(
                f"[DEBUG] 遍历完成: sec_ii_para={sec_ii_para}, sec_iii_para={sec_iii_para}"
            )

            if sec_ii_para is not None:
                RuntimeLogger.log(f"[DEBUG] 开始处理第二节，sec_ii_para={sec_ii_para}")
                # 寻找列表起始位置 "本项目送审分为"
                list_start_idx = -1
                for j in range(sec_ii_para + 1, len(doc.paragraphs)):
                    if "本项目送审分为" in doc.paragraphs[j].text:
                        list_start_idx = j
                        RuntimeLogger.log(f"[DEBUG] 找到'本项目送审分为'在段落 {j}")
                        break
                    if sec_iii_para and j >= sec_iii_para:
                        break

                if list_start_idx != -1:
                    RuntimeLogger.log(
                        f"[DEBUG] list_start_idx={list_start_idx}, 开始处理子项目列表"
                    )
                    # 1. 记录并彻底清理原有占位内容 (标题句之后，三、之前)
                    search_end_idx = (
                        sec_iii_para if sec_iii_para else len(doc.paragraphs)
                    )

                    # 寻找结束位置：包含“评审原则”或到达三、审核结果
                    real_end_idx = search_end_idx
                    for k in range(list_start_idx + 1, search_end_idx):
                        if "评审原则" in doc.paragraphs[k].text:
                            real_end_idx = k
                            break

                    p_title = doc.paragraphs[list_start_idx]
                    subs = context.get("sub_projects", [])
                    sub_count = len(subs) if subs else 1

                    # 强制设置段落间距和缩进的辅助函数
                    def force_paragraph_format(p):
                        pf = p.paragraph_format
                        pf.space_before = Pt(0)
                        pf.space_after = Pt(0)
                        # 【恢复】使用单倍行距
                        pf.line_spacing = None
                        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                        pf.first_line_indent = Pt(22)  # 约两个中文字符缩进

                        # [关键修复] 取消对齐到网格，解决中文字符被撑大的问题
                        ReceiptProcessor._disable_snap_to_grid(p)

                    # 使用阿拉伯数字
                    p_title.text = ""
                    force_paragraph_format(p_title)

                    r_title = p_title.add_run(f"本项目送审分为{sub_count}个子项目：")

                    # 应用字体
                    r_title.font.name = sample_font_name
                    r_title._element.rPr.rFonts.set(qn("w:eastAsia"), sample_font_name)
                    if sample_font_size:
                        r_title.font.size = sample_font_size

                    # [关键修复] 设置在 _element 上，因为 p 对象是动态生成的
                    p_title._element._processed_by_agent = True

                    # 倒序删除标题句与评审原则之间的所有段落
                    # [修复] 预先收集元素，避免索引失效
                    paras_to_del = []
                    for k in range(list_start_idx + 1, real_end_idx):
                        paras_to_del.append(doc.paragraphs[k])

                    for p in reversed(paras_to_del):
                        p_el = p._element
                        p_el.getparent().remove(p_el)

                    # 重新计算插入位置 (标题句之后)
                    # 由于我们删除了中间的所有段落，原本在 real_end_idx 的段落现在紧跟在 p_title 之后
                    # 重新通过 XML 查找位置以确保准确
                    insert_pos = -1
                    for i, p in enumerate(doc.paragraphs):
                        if p._element == p_title._element:
                            insert_pos = i + 1
                            break

                    if insert_pos == -1:
                        insert_pos = list_start_idx + 1

                    subs = context.get("sub_projects", [])
                    # 调试日志
                    RuntimeLogger.log(f"[DEBUG] 从 context 获取到 {len(subs)} 个子项目")
                    if subs:
                        for i, s in enumerate(subs, 1):
                            RuntimeLogger.log(
                                f"[DEBUG]   子项目{i}: {s.get('sub_project_name', 'N/A')} - 送审{s.get('submission_fp', 0)}FP, {s.get('submission_days', 0)}人天"
                            )
                    else:
                        RuntimeLogger.log("[DEBUG] sub_projects 为空，将使用单项目模式")

                    # ... (其后保持 subs 逻辑)
                    if not subs:
                        # 单项目模式
                        sub_name = context.get("sub_project_name", "子项目")
                        if not sub_name or sub_name == "子项目":
                            sub_name = context["project_name"]

                        subs = [
                            {
                                "sub_project_name": sub_name,
                                "submission_fp": context.get("total_fp", "0"),
                                "submission_days": context.get("submission_days", "0"),
                                "new_fp": context.get("new_fp", "0"),
                                "new_ratio": context.get("new_ratio", "0%"),
                                "reuse_fp": context.get("reuse_fp", "0"),
                                "reuse_ratio": context.get("reuse_ratio", "0%"),
                                "legacy_fp": context.get("legacy_fp", "0"),
                                "legacy_ratio": context.get("legacy_ratio", "0%"),
                                "eval_days": context.get("eval_days", "0"),
                            }
                        ]

                    for idx, sub in enumerate(subs, 1):
                        # Line 1: 1、项目名称：送审xxx功能点，送审工作量自评为xxx人天。
                        p1 = doc.paragraphs[insert_pos].insert_paragraph_before()
                        force_paragraph_format(p1)

                        # 重要：送审总功能点评定 = 新增 + 复用 + 利旧
                        try:
                            calc_total_fp = (
                                int(float(str(sub.get("new_fp", 0)).replace(",", "")))
                                + int(
                                    float(str(sub.get("reuse_fp", 0)).replace(",", ""))
                                )
                                + int(
                                    float(str(sub.get("legacy_fp", 0)).replace(",", ""))
                                )
                            )
                        except:
                            calc_total_fp = sub.get("submission_fp", 0)

                        # 恢复无空格格式
                        r_text = f"{idx}、{sub['sub_project_name']}：送审{format_fp(calc_total_fp)}功能点，送审工作量自评为{format_days(sub['submission_days'])}人天。"
                        r1 = p1.add_run(r_text)

                        # 应用字体
                        r1.font.name = sample_font_name
                        r1._element.rPr.rFonts.set(qn("w:eastAsia"), sample_font_name)
                        if sample_font_size:
                            target_size = sample_font_size
                            if not isinstance(target_size, int) and hasattr(
                                target_size, "pt"
                            ):
                                # 如果是 Length 对象，保持原样
                                pass
                            r1.font.size = target_size

                        p1._element._processed_by_agent = True
                        insert_pos += 1

                        # Line 2: 评估结果：新增功能点 2364 个，占比 100.0%；复用功能点 0 个，占比 0.0%；利旧功能点 0 个，占比 0.0%，评定工作量 214.91 人天。
                        p2 = doc.paragraphs[insert_pos].insert_paragraph_before()
                        force_paragraph_format(p2)

                        r_label = p2.add_run("评估结果：")
                        r_label.bold = True

                        # 准备数据，确保百分比不重复 %
                        r_new = str(sub.get("new_ratio", "0%")).replace("%", "") + "%"
                        r_reuse = (
                            str(sub.get("reuse_ratio", "0%")).replace("%", "") + "%"
                        )
                        r_legacy = (
                            str(sub.get("legacy_ratio", "0%")).replace("%", "") + "%"
                        )

                        # 恢复无空格格式，分号/逗号严格遵循模板
                        text2 = f"新增功能点{format_fp(sub['new_fp'])}个，占比{r_new}；复用功能点{format_fp(sub['reuse_fp'])}个，占比{r_reuse}；利旧功能点{format_fp(sub['legacy_fp'])}个，占比{r_legacy}，评定工作量{format_days(sub['eval_days'])}人天。"
                        r_val = p2.add_run(text2)

                        # 应用字体
                        for r in [r_label, r_val]:
                            r.font.name = sample_font_name
                            r._element.rPr.rFonts.set(
                                qn("w:eastAsia"), sample_font_name
                            )
                            if sample_font_size:
                                target_size = sample_font_size
                                if not isinstance(target_size, int) and hasattr(
                                    target_size, "pt"
                                ):
                                    pass
                                r.font.size = target_size

                        p2._element._processed_by_agent = True
                        insert_pos += 1

                    # 清除结束 (旧的 delete_targets 逻辑已上移至插入前)
                    pass
                    pass

        def handle_section_iii_table():
            """处理第三节的项目明细表：如果有多子项目则增加行"""
            if not context.get("is_merge") or not context.get("sub_projects"):
                return

            # 寻找第三节下方的第一个表格
            sec_iii_idx = -1
            for i, p in enumerate(doc.paragraphs):
                if "三、审核结果" in p.text.replace(" ", ""):
                    sec_iii_idx = i
                    break

            if sec_iii_idx == -1:
                return

            # 在三、审核结果之后的第一个表格
            target_table = None
            for table in doc.tables:
                # 检查表格是否在三、审核结果之后
                # (doc.tables 不存储段落索引，我们需要通过 _element 查找)
                if table._element.getparent().index(table._element) > doc.paragraphs[
                    sec_iii_idx
                ]._element.getparent().index(doc.paragraphs[sec_iii_idx]._element):
                    target_table = table
                    break

            if not target_table:
                return

            sub_projects = context["sub_projects"]
            if len(sub_projects) <= 1:
                return

            # 寻找包含“项目名称”且有高亮占位符的样板行
            template_row_idx = -1
            for i, row in enumerate(target_table.rows):
                row_text = "".join(cell.text for cell in row.cells)
                # 如果这一行有 X 或者 混合了项目名称标签，则认为是模板行
                if (
                    any(cell.text.strip() in ["X", "x"] for cell in row.cells)
                    or "项目名称" in row_text
                ):
                    template_row_idx = i
                    break

            if template_row_idx == -1:
                # 默认最后一行（排除合计行的情况，通常合计行会有“合计”字样）
                template_row_idx = len(target_table.rows) - 1
                for i in range(len(target_table.rows) - 1, -1, -1):
                    if "合计" not in "".join(
                        cell.text for cell in target_table.rows[i].cells
                    ):
                        template_row_idx = i
                        break

            # 保存模板行的格式
            template_row = target_table.rows[template_row_idx]

            # 为剩余的子项目增加行 (并填入内容)
            for i in range(1, len(sub_projects)):
                sub = sub_projects[i]
                new_row = target_table.add_row()

                # 准备该子项目的数据包
                try:
                    calc_total = (
                        int(float(str(sub.get("new_fp", 0)).replace(",", "")))
                        + int(float(str(sub.get("reuse_fp", 0)).replace(",", "")))
                        + int(float(str(sub.get("legacy_fp", 0)).replace(",", "")))
                    )
                except:
                    calc_total = sub.get("submission_fp", 0)

                sub_bundle = {
                    "total_fp": format_fp(calc_total),
                    "new_fp": format_fp(sub.get("new_fp", 0)),
                    "new_ratio": str(sub.get("new_ratio", "0%")),
                    "reuse_fp": format_fp(sub.get("reuse_fp", 0)),
                    "reuse_ratio": str(sub.get("reuse_ratio", "0%")),
                    "legacy_fp": format_fp(sub.get("legacy_fp", 0)),
                    "legacy_ratio": str(sub.get("legacy_ratio", "0%")),
                    "submission_days": format_days(sub.get("submission_days", "0")),
                    "eval_days": format_days(sub.get("eval_days", "0")),
                }

                for j, cell in enumerate(template_row.cells):
                    # 复制文本内容
                    cell_text = template_row.cells[j].text
                    new_row.cells[j].text = cell_text

                    # 应用该项目的数据
                    if new_row.cells[j].paragraphs:
                        p = new_row.cells[j].paragraphs[0]
                        p.alignment = template_row.cells[j].paragraphs[0].alignment
                        # 标记已处理，防止外层循环重复处理
                        ReceiptProcessor._smart_replace_highlights(p, sub_bundle)
                        ReceiptProcessor._disable_snap_to_grid(p)
                        # [关键] 直接在 _element 上标记
                        p._element._processed_by_agent = True

            # 最后处理样板行本身 (使用第一个子项目的数据)
            sub0 = sub_projects[0]
            try:
                calc_total0 = (
                    int(float(str(sub0.get("new_fp", 0)).replace(",", "")))
                    + int(float(str(sub0.get("reuse_fp", 0)).replace(",", "")))
                    + int(float(str(sub0.get("legacy_fp", 0)).replace(",", "")))
                )
            except:
                calc_total0 = sub0.get("submission_fp", 0)

            sub0_bundle = {
                "total_fp": format_fp(calc_total0),
                "new_fp": format_fp(sub0.get("new_fp", 0)),
                "new_ratio": str(sub0.get("new_ratio", "0%")),
                "reuse_fp": format_fp(sub0.get("reuse_fp", 0)),
                "reuse_ratio": str(sub0.get("reuse_ratio", "0%")),
                "legacy_fp": format_fp(sub0.get("legacy_fp", 0)),
                "legacy_ratio": str(sub0.get("legacy_ratio", "0%")),
                "submission_days": format_days(sub0.get("submission_days", "0")),
                "eval_days": format_days(sub0.get("eval_days", "0")),
            }
            for cell in template_row.cells:
                if cell.paragraphs:
                    p = cell.paragraphs[0]
                    ReceiptProcessor._smart_replace_highlights(p, sub0_bundle)
                    ReceiptProcessor._disable_snap_to_grid(p)
                    # [关键] 直接在 _element 上标记
                    p._element._processed_by_agent = True

        handle_section_ii()
        handle_section_iii_table()

        # 2.5 调试：检查标记是否有效
        # marked_count = sum(1 for p in doc.paragraphs if hasattr(p._element, "_processed_by_agent") and p._element._processed_by_agent)
        # print(f"DEBUG: {marked_count} 个段落已标记为 _processed_by_agent")

        # 3. 准备 Section III (审核结果) 基础替换词
        reduction_val = context.get("reduction_ratio", "0%")

        # 处理送审时间 (从 yyyy-MM-dd 转为 中文年月日)
        date_str = context["submission_time"]
        try:
            dt = datetime.strptime(date_str, "%Y-%m-%d")
            cn_date = (
                dt.strftime("%Y年%m月%d日").replace("年0", "年").replace("月0", "月")
            )
        except:
            cn_date = date_str

        # 准备基础替换词
        text_replacements = {
            "2025 年 x 月 x 日": cn_date,
            "2025年x月x日": cn_date,
            "2025年X月X日": cn_date,
            "2026 年 x 月 x 日": cn_date,
            "2026年x月x日": cn_date,
            "2026年X月X日": cn_date,
            "XXX人天": data_bundle["eval_days"],
            "xxxx人天": data_bundle["eval_days"],
            "xxxx 人天": data_bundle["eval_days"],
            "xxxxxxxx人天": data_bundle["submission_days"],
            "xxxxxxxx 人天": data_bundle["submission_days"],
            "xxxx功能点": data_bundle["total_fp"],
            "xxxx 功能点": data_bundle["total_fp"],
            "xx.xx%": reduction_val,
            "XXXXX%": reduction_val,
            "XXXX%": reduction_val,
            "XXX%": reduction_val,
            "xxxxx%": reduction_val,
            "xxxx%": reduction_val,
            "xxx%": reduction_val,
        }

        # 动态补充项目名称占位符 (XX, XXX, XXXX, ... XXXXXXXX)
        # 注意：不再添加纯X占位符，统一由Section C的正则处理
        # for i in range(10, 1, -1):
        #     # 大写 X
        #     ph_upper = "X" * i
        #     if ph_upper not in text_replacements:
        #         text_replacements[ph_upper] = context["project_name"]
        #     # 小写 x
        #     ph_lower = "x" * i
        #     if ph_lower not in text_replacements:
        #         text_replacements[ph_lower] = context["project_name"]
        #     # 混合/带数字
        #     if i <= 5:
        #         text_replacements[f"XX{i}"] = context["project_name"]

        # 增加一些特殊变体
        # text_replacements["Xx"] = context["project_name"]
        # text_replacements["Xxx"] = context["project_name"]

        # 定义标签追加 (如 项目名称：后追加)
        labels = {
            "项目名称：": context["project_name"],
            "项目名称:": context["project_name"],
            "项目名称": context["project_name"],
            "项目编号：": context["project_id"],
            "项目编号:": context["project_id"],
            "项目编号": context["project_id"],
            "送审单位：": context["submission_unit"],
            "送审单位:": context["submission_unit"],
            "送审单位": context["submission_unit"],
            "送审人：": context["submitter"],
            "送审人:": context["submitter"],
            "送审人": context["submitter"],
            "项目负责人：": context["submitter"],
            "项目负责人:": context["submitter"],
            "项目负责人": context["submitter"],
            "送审时间：": context["submission_time"],
            "送审时间:": context["submission_time"],
            "送审时间": context["submission_time"],
            "日期：": context["submission_time"],
            "日期": context["submission_time"],
        }

        # 定义一个统一的处理器
        def process_paragraph(p, extra_context=""):
            # [关键修复] 取消对齐到网格，解决中文字符被撑大的问题
            ReceiptProcessor._disable_snap_to_grid(p)

            # [关键修复] 使用 _element 判断，因为 Paragraph 对象是动态创建的
            if (
                hasattr(p._element, "_processed_by_agent")
                and p._element._processed_by_agent
            ):
                return

            # A. 处理高亮色块
            ReceiptProcessor._smart_replace_highlights(p, data_bundle, extra_context)

            # B. 处理特定文本替换 (日期、百分比等) - 移动到 X 模式之前，并支持跨 Run 替换
            p_text = p.text
            sorted_replaces = sorted(
                text_replacements.items(), key=lambda x: len(x[0]), reverse=True
            )
            replaced_any = False
            for old, new in sorted_replaces:
                if old in p_text:
                    # 尝试直接在 Run 中替换
                    found_in_run = False
                    for run in p.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, str(new))
                            found_in_run = True

                    # 如果 Run 中没找到（说明被碎裂了），则全段替换并尝试保留样式
                    if not found_in_run:
                        p_text = p_text.replace(old, str(new))
                        replaced_any = True

            if replaced_any:
                # 获取原始样式并重组段落
                sample_run = p.runs[0] if p.runs else None
                for r in p.runs:
                    if not sample_run and r.text.strip():
                        sample_run = r
                    r.text = ""
                if p.runs:
                    p.runs[0].text = p_text
                    target_run = p.runs[0]
                else:
                    target_run = p.add_run(p_text)

                # 恢复样板样式
                if sample_run:
                    if sample_run.font.name:
                        target_run.font.name = sample_run.font.name
                        try:
                            target_run._element.rPr.rFonts.set(
                                qn("w:eastAsia"), sample_run.font.name
                            )
                        except:
                            pass
                    if sample_run.font.size:
                        target_run.font.size = sample_run.font.size
                    if sample_run.font.bold is not None:
                        target_run.font.bold = sample_run.font.bold
                    if sample_run.font.color and sample_run.font.color.rgb:
                        try:
                            target_run.font.color.rgb = sample_run.font.color.rgb
                        except:
                            pass

                # 标记已处理并退出，避免后续 X 替换和标签追加再次处理
                p._element._processed_by_agent = True
                return

            # C. 统一占位符替换 (使用正则处理所有 X/x 连缀)
            # 优化正则：匹配独立的 X 或 多个 X 连缀，且不影响其它英文单词（如Excel）
            # 排除后面紧跟 % 的情况，避免误杀比例
            # 使用负回顾断言 (?<![a-zA-Z]) 替代变长回顾断言 (?<=^|[^a-zA-Z]) 以适配 Python re
            x_pattern = re.compile(
                r"(?i)(?<![a-zA-Z])[Xx]{1,}(?=[^a-zA-Z%]|$)|(?i)[Xx]{2,}"
            )

            p_text = p.text
            # 检查是否包含 x 或 X
            if any(c in p_text for c in "xX") and x_pattern.search(p_text):
                # 寻找样板 Run (用于继承样式)
                sample_run = p.runs[0] if p.runs else None
                for r in p.runs:
                    if any(c in r.text for c in "xX"):
                        sample_run = r
                        break

                # 执行正则替换：根据不同模式替换为项目名称
                def replace_x_pattern(match):
                    # 【优化】智能处理：避免出现“项目项目”连读
                    proj_name = context["project_name"]
                    p_full_text = p.text
                    match_end = match.end()

                    # 检查占位符后面是否紧跟“项目”二字
                    after_text = p_full_text[match_end:].strip()
                    if after_text.startswith("项目") and proj_name.endswith("项目"):
                        return proj_name[:-2]  # 返回去掉“项目”后缀的名称
                    return proj_name

                new_text = x_pattern.sub(replace_x_pattern, p_text)

                if new_text != p_text:
                    # 获取原始样式
                    fname = sample_run.font.name if sample_run else None
                    fsize = sample_run.font.size if sample_run else None
                    fbold = sample_run.font.bold if sample_run else None
                    fcolor = (
                        sample_run.font.color.rgb
                        if (sample_run and sample_run.font.color)
                        else None
                    )

                    # 重组段落 (防止碎裂同时彻底清除旧 placeholder)
                    for r in p.runs:
                        r.text = ""
                    if p.runs:
                        p.runs[0].text = new_text
                        target_run = p.runs[0]
                    else:
                        target_run = p.add_run(new_text)

                    # 恢复样式
                    if fname:
                        target_run.font.name = fname
                        try:
                            target_run._element.rPr.rFonts.set(qn("w:eastAsia"), fname)
                        except:
                            pass
                    if fsize:
                        target_run.font.size = fsize
                    if fbold is not None:
                        target_run.font.bold = fbold
                    if fcolor:
                        try:
                            target_run.font.color.rgb = fcolor
                        except:
                            pass

                    # 只要处理了 X 占位符，就标记已完成，防止后续标签逻辑再次追加
                    p._processed_by_agent = True
                    return

            # D. 处理标签追加 (项目名称：[空/弱值])
            p_text_clean = "".join(p.text.split()).replace(" ", "")
            for label, val in labels.items():
                label_clean = "".join(label.split()).replace(" ", "")
                if p_text_clean.startswith(label_clean):
                    # 检查标签后面是否是“弱值”
                    suffix = p_text_clean[len(label_clean) :].strip()
                    weak_base = [
                        "",
                        "N/A",
                        "[]",
                        "【】",
                        "：",
                        ":",
                        "1",
                        "x",
                        "X",
                        "unknown",
                    ]
                    is_weak = (
                        not suffix
                        or suffix in weak_base
                        or all(c in "_.-·— " for c in suffix)
                    )

                    if is_weak:
                        if val and str(val) != "N/A":
                            # 清理弱值
                            if suffix and suffix not in ["：", ":"]:
                                for r in p.runs:
                                    if suffix in r.text:
                                        r.text = r.text.replace(suffix, "", 1)

                            if not p.text.strip().endswith(("：", ":")):
                                p.add_run("：")

                            new_run = p.add_run(f"{val}")
                            # 继承该段落之前的样式
                            ref = p.runs[0] if p.runs else None
                            if ref:
                                new_run.font.name = ref.font.name
                                new_run.font.size = ref.font.size
                                new_run.font.bold = ref.font.bold
                                try:
                                    new_run._element.rPr.rFonts.set(
                                        qn("w:eastAsia"), ref.font.name
                                    )
                                except:
                                    pass
                        p._element._processed_by_agent = True
                        return

            p._element._processed_by_agent = True

        # 3. 遍历所有段落
        for p in doc.paragraphs:
            process_paragraph(p)

        # 4. 遍历所有表格 (使用列头信息处理高亮块和普通文本)
        for table in doc.tables:
            # 记录表头
            headers = []
            if len(table.rows) > 0:
                for cell in table.rows[0].cells:
                    headers.append(cell.text.strip())

            # 【优化】如果当前表看起来像是明细表，且有多个子项目
            is_detail_table = False
            if context.get("is_merge") and context.get("sub_projects"):
                row_contents = ["".join(c.text for c in r.cells) for r in table.rows]
                if any("项目名称" in rc for rc in row_contents) or any(
                    "送审工作量" in rc for rc in row_contents
                ):
                    is_detail_table = True

            sub_idx = 0
            for row_idx, row in enumerate(table.rows):
                # 确定当前行应该使用哪个数据环境
                row_bundle = data_bundle
                if is_detail_table:
                    # 如果不是表头行且不是合计行
                    row_text = "".join(cell.text for cell in row.cells)
                    if row_idx > 0 and "合计" not in row_text:
                        subs = context.get("sub_projects", [])
                        if sub_idx < len(subs):
                            sub = subs[sub_idx]
                            # 构建特定行的 bundle
                            row_bundle = {
                                "new_fp": sub.get("new_fp", 0),
                                "new_ratio": sub.get("new_ratio", "0%"),
                                "reuse_fp": sub.get("reuse_fp", 0),
                                "reuse_ratio": sub.get("reuse_ratio", "0%"),
                                "legacy_fp": sub.get("legacy_fp", 0),
                                "legacy_ratio": sub.get("legacy_ratio", "0%"),
                                "submission_days": sub.get("submission_days", "0"),
                                "eval_days": sub.get("eval_days", "0"),
                                "total_fp": sub.get("submission_fp", 0),
                            }
                            # 增加对单元格内项目名称的单独处理
                            for cell in row.cells:
                                if any(x in cell.text for x in ["X", "x", "子项目"]):
                                    # 针对项目名称格，尝试直接填充
                                    if (
                                        "名称" in headers[row.cells.index(cell)]
                                        or "子项目" in headers[row.cells.index(cell)]
                                    ):
                                        cell.text = sub["sub_project_name"]
                            sub_idx += 1

                for idx, cell in enumerate(row.cells):
                    # 获取列上下文 (表头)
                    col_header = headers[idx] if idx < len(headers) else ""
                    for p in cell.paragraphs:
                        # 使用当前行的 bundle
                        ReceiptProcessor._smart_replace_highlights(
                            p, row_bundle, col_header
                        )
                        process_paragraph(p, col_header)

    @staticmethod
    def _smart_replace_highlights(p, data_bundle, col_header=""):
        """
        根据"标点隔离+权重路径+全句环视"原则精准填报。
        改进版：将连续的高亮 runs 识别为一个高亮块，每个块只填充一次。
        """
        runs = p.runs
        header = col_header.strip().replace(" ", "")
        major_seps = ["。", "；", "！", "？", "\n", "\r", "\t"]
        minor_seps = ["，", "：", ":", " ", ","]
        all_seps = major_seps + minor_seps

        # 第一步：识别所有高亮块（连续的高亮 runs 视为一个块）
        highlight_blocks = []
        current_block = []

        for i, run in enumerate(runs):
            # 兼容性修复：某些 docx 文件的高亮值为 'none' 会导致 python-docx 抛出 ValueError
            try:
                has_highlight = run.font.highlight_color is not None
            except (ValueError, Exception):
                has_highlight = False

            if has_highlight:
                current_block.append(i)
            else:
                if current_block:
                    highlight_blocks.append(current_block)
                    current_block = []

        # 最后一个块
        if current_block:
            highlight_blocks.append(current_block)

        # 第二步：对每个高亮块，只填充一次
        for block_indices in highlight_blocks:
            if not block_indices:
                continue

            # 使用第一个 run 的索引进行上下文分析
            i = block_indices[0]

            try:
                # 1. 动态提取本句内的短上下文 (隔离)
                lc = ""  # 左侧短语境
                for j in range(i - 1, -1, -1):
                    txt = runs[j].text
                    has_sep = False
                    for s in all_seps:
                        if s in txt:
                            lc = txt[txt.rfind(s) + 1 :] + lc
                            has_sep = True
                            break
                    if has_sep:
                        break
                    lc = txt + lc

                # 右侧短语境 - 从高亮块结束后开始
                rc = ""
                last_idx = block_indices[-1]
                for j in range(last_idx + 1, len(runs)):
                    txt = runs[j].text
                    has_sep = False
                    for s in all_seps:
                        if s in txt:
                            rc += txt[: txt.find(s)]
                            has_sep = True
                            break
                    if has_sep:
                        break
                    rc += txt

                lc_clean = lc.strip().replace(" ", "")
                rc_clean = rc.strip().replace(" ", "")

                # 合并整个高亮块的文本用于判定
                own_text = (
                    "".join([runs[idx].text for idx in block_indices])
                    .strip()
                    .replace(" ", "")
                )

                # 2. 维度判定 (Dimension) - 优先级：右侧 > 自身 > 左侧 > 表头
                dim = "count"
                if (
                    "%" in rc_clean
                    or "%" in own_text
                    or "占比" in lc_clean
                    or "比例" in lc_clean
                    or "比例" in header
                ):
                    dim = "ratio"
                elif (
                    "人天" in rc_clean
                    or "工作量" in rc_clean
                    or "人天" in lc_clean
                    or "工作量" in lc_clean
                    or "人天" in header
                ):
                    dim = "days"
                elif "个" in rc_clean or "功能点" in rc_clean or "功能点" in header:
                    dim = "count"

                # 3. 对象判定 (Object Scoring)
                obj_map = {
                    "核减": "reduction",
                    "新增": "new",
                    "增加": "new",
                    "复用": "reuse",
                    "优化": "reuse",
                    "利旧": "legacy",
                    "送审": "submission",
                    "核定": "eval",
                    "核准": "eval",
                    "评估": "eval",
                    "评定": "eval",
                }

                obj = None
                # A. 优先判定后置语义 (rc 中的"计新增"等)
                if "计新增" in rc_clean:
                    obj = "new"
                    dim = "ratio"
                elif "计复用" in rc_clean:
                    obj = "reuse"
                    dim = "ratio"
                elif "计利旧" in rc_clean:
                    obj = "legacy"
                    dim = "ratio"

                if obj is None:
                    # B. 寻找离色块最近的操作对象 (全句追溯，不停于逗号)
                    # 我们向上回溯直到大标点（句号/分号）
                    long_lc = ""
                    for j in range(i - 1, -1, -1):
                        txt = runs[j].text
                        if any(s in txt for s in major_seps):
                            parts = [txt.rfind(s) for s in major_seps if s in txt]
                            long_lc = txt[max(parts) + 1 :] + long_lc
                            break
                        long_lc = txt + long_lc

                    long_lc = long_lc.replace(" ", "")
                    last_pos = -1
                    for kw, target_obj in obj_map.items():
                        p_idx = long_lc.rfind(kw)
                        if p_idx > last_pos:
                            last_pos = p_idx
                            obj = target_obj

                    # C. 兜底看表头
                    if obj is None:
                        for kw, target_obj in obj_map.items():
                            if kw in header:
                                obj = target_obj
                                break

                # 4. 映射数据
                val = None
                if obj == "new":
                    val = (
                        data_bundle["new_ratio"]
                        if dim == "ratio"
                        else data_bundle["new_fp"]
                    )
                elif obj == "reuse":
                    val = (
                        data_bundle["reuse_ratio"]
                        if dim == "ratio"
                        else data_bundle["reuse_fp"]
                    )
                elif obj == "legacy":
                    val = (
                        data_bundle["legacy_ratio"]
                        if dim == "ratio"
                        else data_bundle["legacy_fp"]
                    )
                elif obj == "reduction":
                    val = data_bundle["reduction_ratio"]
                elif obj == "submission":
                    val = (
                        data_bundle["submission_days"]
                        if dim == "days"
                        else data_bundle["total_fp"]
                    )
                elif obj == "eval":
                    val = (
                        data_bundle["eval_days"]
                        if dim == "days"
                        else data_bundle["total_fp"]
                    )

                # 5. 执行替换 - 只在第一个 run 中填充，其他 run 清空
                if val is not None:
                    val_str = str(val)
                    # 清理重复的百分号
                    if "%" in val_str and ("%" in rc_clean or own_text.endswith("%")):
                        val_str = val_str.replace("%", "")

                    # 第一个 run 填充值
                    runs[block_indices[0]].text = val_str
                    runs[block_indices[0]].font.highlight_color = None

                    # 其他 run 清空（删除多余的 X 占位符）
                    for idx in block_indices[1:]:
                        runs[idx].text = ""
                        runs[idx].font.highlight_color = None
            except Exception:
                pass


def datetime_str():
    return datetime.now().strftime("%Y%m%d_%H%M%S")
